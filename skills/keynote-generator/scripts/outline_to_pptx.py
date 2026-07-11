#!/usr/bin/env python3
"""Create a verifiable PPTX fallback from the branded outline JSON.

This does not replace the native `.key` generation on macOS. It exists to verify
branding rules in non-macOS environments:
- content slides: white background, black text, logo watermark top-right
- chapter slides: yellow #ffed3a background, black text, no watermark
- speaker notes on every slide
"""
from __future__ import annotations

import base64
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

YELLOW = RGBColor(0xFF, 0xED, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
SKILL_ROOT = Path(__file__).resolve().parents[1]
LOGO_B64 = SKILL_ROOT / "assets" / "aulab_logo_watermark_48.png.b64"
DEFAULT_LOGO_PATH = Path("/workspace/work/keynote-generator-assets/aulab_logo_watermark_48.png")


def materialize_logo_asset() -> Path:
    if DEFAULT_LOGO_PATH.exists() and DEFAULT_LOGO_PATH.read_bytes()[:8].hex() == "89504e470d0a1a0a":
        return DEFAULT_LOGO_PATH
    if not LOGO_B64.exists():
        raise FileNotFoundError(f"Missing skill logo asset: {LOGO_B64}")
    raw = base64.b64decode("".join(LOGO_B64.read_text(encoding="ascii").split()))
    DEFAULT_LOGO_PATH.parent.mkdir(parents=True, exist_ok=True)
    DEFAULT_LOGO_PATH.write_bytes(raw)
    if DEFAULT_LOGO_PATH.read_bytes()[:8].hex() != "89504e470d0a1a0a":
        raise ValueError(f"Materialized logo is not a valid PNG: {DEFAULT_LOGO_PATH}")
    return DEFAULT_LOGO_PATH


def set_background(slide: Any, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide: Any, left: float, top: float, width: float, height: float, text: str, size: int, bold: bool = False, align: Any | None = None) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    return box


def add_bullets(slide: Any, bullets: list[str]) -> None:
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.0), Inches(4.6))
    frame = box.text_frame
    frame.clear()
    for idx, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(28)
        p.font.color.rgb = BLACK


def add_logo(slide: Any, logo_path: str) -> None:
    path = Path(logo_path)
    if not path.exists():
        path = materialize_logo_asset()
    if path.read_bytes()[:8].hex() != "89504e470d0a1a0a":
        raise ValueError(f"Logo is not a valid PNG: {path}")
    slide.shapes.add_picture(str(path), Inches(12.55), Inches(0.18), width=Inches(0.45))


def add_notes(slide: Any, notes: list[str]) -> None:
    notes_text = "\n".join(notes).strip() or "Presentare il contenuto della slide."
    slide.notes_slide.notes_text_frame.text = notes_text


def build_pptx(outline: dict[str, Any], output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    logo_path = outline.get("metadata", {}).get("logo_path") or str(materialize_logo_asset())

    for item in outline.get("slides", []):
        slide = prs.slides.add_slide(blank_layout)
        slide_type = item.get("type", "content")
        title = item.get("title", "Untitled")
        bullets = [str(b) for b in item.get("bullets", [])]
        notes = [str(n) for n in item.get("notes", [])]

        if slide_type == "chapter":
            set_background(slide, YELLOW)
            add_textbox(slide, 1.0, 2.85, 11.3, 1.2, title, 46, bold=True, align=PP_ALIGN.CENTER)
        else:
            set_background(slide, WHITE)
            add_textbox(slide, 0.75, 0.55, 11.2, 0.8, title, 34, bold=True)
            add_bullets(slide, bullets)
            if item.get("watermark", True):
                add_logo(slide, logo_path)
        add_notes(slide, notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("Usage: outline_to_pptx.py outline.json output.pptx", file=sys.stderr)
        return 2
    outline_path = Path(argv[1])
    output_path = Path(argv[2])
    outline = json.loads(outline_path.read_text(encoding="utf-8"))
    build_pptx(outline, output_path)
    slides = outline.get("slides", [])
    chapters = sum(1 for s in slides if s.get("type") == "chapter")
    print(f"pptx_created= {output_path}")
    print(f"slides= {len(slides)}")
    print(f"chapters= {chapters}")
    print(f"logo_path= {outline.get('metadata', {}).get('logo_path') or DEFAULT_LOGO_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
